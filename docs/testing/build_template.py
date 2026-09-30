"""Build the focused four-demo fill-in template. Requires python-docx."""
from pathlib import Path
import subprocess
import json
from datetime import datetime, timezone
ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).resolve().parent
revision=subprocess.check_output(['git','rev-parse','--short','HEAD'],cwd=ROOT,text=True).strip()
sections=[]
def group(title,intro=''):
    g={'title':title,'intro':intro,'cases':[]};sections.append(g);return g
def case(g,id,description,files,steps,expected):
    g['cases'].append(dict(id=id,description=description,files=files,steps=steps,expected=expected))
setup="""Scope: ONLY the four assigned combinations: short text → PNG, longer text → WAV, WAV payload → PNG, and PNG payload → WAV. There are four complete positive flows and twelve focused validation/negative/comparison checks. No video, full format matrix, API-only tests, or exhaustive depth sweep is required by this document.

Tester __________; Date __________; Browser/OS __________; App commit __________; Alice public-key fingerprint __________; overall Pass / Fail / Blocked __________.

Run python backend/app.py in your existing project environment; open http://127.0.0.1:5000. Keep the server running. Use SHA-256, Alice's existing signing key, manual offset 100 unless a row changes it. The backend uses its saved Bob encryption key automatically. Keep both key pairs unchanged during the test run.

Save each demo in its own folder, e.g. test_evidence/manual/D1/; the app reuses stego_image.png and stego_audio.wav filenames. Export Alice's public key once. Download and upload the saved stego file during Verify rather than relying on the shortcut that copies settings.

Run D1–D4 first and keep the four Authentic baseline files. After each negative check, restore its original file, matching settings and trusted public key. Alter hidden content / Alter cover value use the LAST CREATED protected file: recreate the specified baseline immediately before clicking either button. Do not assume those buttons modify an arbitrary received upload.

Fill Observed result, Pass/Fail, and screenshot/report/defect ID. Mark phases not applicable to a row as N/A. Browser validation wording can vary; record the actual wording. Extra controls/rows are not a requirement to test every file in /samples.
"""
behavior="""Authentic requires extraction, decryption, record binding, signature, payload integrity and stable cover integrity all Passed; reason_code ALL_CHECKS_PASSED. Content must be recovered exactly. Stable cover hashing masks the selected low bits, so it permits legitimate LSB embedding changes.

Important negative-case distinction: a wrong Alice public key produces Signature Invalid; changing a protected cover bit with manual extraction produces Tampered / COVER_HASH_MISMATCH; changing the encrypted hidden package produces Cannot Verify / DECRYPTION_FAILED. Wrong extraction depth or location normally produces Payload Missing; it does not prove the file was tampered with.

Current implementation can show/download parsed content after Signature Invalid or Tampered. Treat that content as untrusted. Decryption failure should release no content. A Compare success is a measurement result, not an authenticity verdict.

Fit depends on the whole protected package, including metadata, hashes, signature, encryption headers and padding. Basic capacity = floor((N−start)×LSBs/8), but the exact required carrier units shown in Review decide fit. The app does not silently increase your selected depth. All expectations assume the normal saved RSA2048 keys and modest default metadata.
"""
workflow="""For D1–D4, follow all three phases:

Protect: choose cover; Text message or File; enter/select payload; SHA-256; specified LSB depth; manual offset 100. Review capacity; record content/package/available bytes and Fit. If it fits, Sign & create stego file; download it to the demo folder.

Verify: upload that saved stego PNG/WAV; enter the same depth and offset; import Alice's matching public key; Extract & verify. Record Authentic, ALL_CHECKS_PASSED, all six Passed checks, and expected/recomputed hash matches. Show recovered text/audio/image. Download recovered file payloads and check exact original bytes/hash if possible. Export the verification report.

Compare: upload the original cover and its matching stego output; Compare objects. Images show histograms and difference views; audio shows aligned waveforms, channel metrics and SNR. Record changed percentage, MSE, maximum difference and PSNR/SNR. Legitimate maximum value change is ≤2^LSBs−1. Tiny differences may be difficult to see or hear. Independently uploaded pairs correctly label hidden content unknown; comparison cannot extract it.

Integrated observations: Protect fit ____ / package ____ B / available ____ B; Verify verdict ____ / reason ____ / checks ____ / recovered content match ____; Compare changed% ____ / MSE ____ / max difference ____ / PSNR ____ / SNR ____; evidence ____; Pass/Fail ____.
"""
inventory=[
('Image cover','samples/covers/astronaut.png','512×512 RGB; 786432 channel units; depth 1/index 100 gives 98291 B capacity.'),
('Audio cover','samples/audio/sample-006.wav','Mono PCM16, 16000 samples at 16000 Hz; capacities at index 100: depth 1=1987 B; depth 2=3975 B; depth 4=7950 B.'),
('WAV payload','samples/audio/sample-002.wav','24044 B; mono PCM16, 12000 samples at 12000 Hz. Fits astronaut PNG at depth 1 with default overhead.'),
('PNG payload / small cover','samples/images/sample-001.png','2650 B; 176×152 RGB; 80256 channel units. Payload fits sample-006 WAV at depth 4; as cover it cannot hold the 24044 B WAV at depth 1.'),
('Optional longer typed text','samples/text/sample-019.txt','If preferred, paste this document into Message instead of Long example. Keep Text message mode.'),
('Wrong original image','samples/covers/chelsea.png','451×300 RGB; intentional dimension mismatch against astronaut.'),
('Wrong original audio','samples/audio/sample-002.wav','Use as wrong original for a sample-006-based stego: rate/sample count differ; deliberate Compare error.'),
]
g=group('A. Your four complete positive demonstrations','Run Protect → Verify → Compare for each row. D1–D4 are the only payload/cover combinations in scope.')
case(g,'D1','Short text payload → PNG cover','Cover: samples/covers/astronaut.png; payload: Short example button in Text message mode.','Depth 1; manual offset 100; complete all three phases.','Protect fits; output stego_image.png. Verify Authentic and exact short text. Image Compare succeeds; max difference ≤1; histograms/difference views render.')
case(g,'D2','Longer text payload → WAV cover','Cover: samples/audio/sample-006.wav; payload: Long example button (or paste samples/text/sample-019.txt).','Depth 2; manual offset 100; complete all three phases.','Protected package fits within 3975 B with defaults; output stego_audio.wav. Verify Authentic and full message preserved. Audio Compare succeeds; max sample difference ≤3; waveform/SNR evidence.')
case(g,'D3','WAV audio payload → PNG cover','Cover: samples/covers/astronaut.png; File payload: samples/audio/sample-002.wav.','Depth 1; manual offset 100; review 24044 B payload plus overhead before creation; complete three phases.','Fits within 98291 B; PNG output; Verify Authentic and exact recovered WAV, audio preview/download available. Image Compare succeeds; max difference ≤1.')
case(g,'D4','PNG image payload → WAV cover','Cover: samples/audio/sample-006.wav; File payload: samples/images/sample-001.png.','Depth 4; manual offset 100; complete all three phases.','2650 B PNG plus overhead fits within 7950 B; WAV output; Verify Authentic and exact recovered PNG with image preview/download. Audio Compare succeeds; max difference ≤15.')
g=group('B. Four focused Protect/settings checks','These reuse the same four combinations. Save alternate valid outputs separately; do not overwrite the baseline files.')
case(g,'P1','D3: cover too small, then sufficient depth','Cover: samples/images/sample-001.png; WAV payload: samples/audio/sample-002.wav.','Review depth 1/index 100; then depth 3/index 100. Create and Verify the depth-3 output if it fits.','Depth 1: no fit (10019 B capacity <24044 B payload even before overhead); create disabled. Depth 3: 30058 B capacity, expected fit with defaults; selected depth retained; Verify Authentic at depth 3.')
case(g,'P2','D4: insufficient LSB capacity, then correct depth','Original D4 cover/payload.','Review at depth 1/index 100; then depth 4.','Depth 1 fails (1987 B <2650 B payload); create disabled/no silent depth change. Depth 4 fits; Verify with depth 4 returns Authentic.')
case(g,'P3','D1: alternate LSB depth and start index','Original D1 cover/short text.','Protect at depth 2/index 200; Review/create; manually Verify depth 2/index 200; Compare.','Fits, retains 2 LSBs and index 200; Authentic; Compare max difference ≤3. Demonstrates offset and depth must agree on both sides.')
case(g,'P4','D1: end-of-cover and out-of-range start','Original D1 cover/short text, depth 1.','Review index 786431 (last valid channel), then 786432 (outside carrier). Restore 100 afterward.','786431: valid index but insufficient capacity/no fit. 786432: outside-cover-range error (0..786431). Neither creates a new stego object; no crash.')
g=group('C. Five focused Verify negative checks','Use saved baselines and change only one variable. Each test must produce a non-Authentic result. Restore correct settings/key/file afterward.')
case(g,'V1','D2: wrong extraction depth','Saved original D2 stego WAV; matching Alice public key.','Verify depth 1/index 100 instead of encoded depth 2. Then restore depth 2.','Wrong depth: normally Payload Missing / NO_RECOGNISED_PAYLOAD; accidental recognizable marker could lead Cannot Verify, but not Authentic. Restored depth 2: Authentic.')
case(g,'V2','D3: wrong extraction start','Saved original D3 stego PNG; matching public key.','Verify depth 1/index 101 or 1000 instead of 100. Then restore 100.','Wrong offset normally Payload Missing / NO_RECOGNISED_PAYLOAD; marker cannot be read. Restore 100: Authentic. Do not expect a distinct Wrong Start Location verdict.')
case(g,'V3','D4: wrong Alice public key','Saved original D4 stego WAV.','Correct depth 4/index 100; click Use wrong key (negative case); verify. Reload matching public key and verify again.','Wrong key: Signature Invalid / SIGNATURE_INVALID; decryption/hashes may still pass; any displayed recovered image is untrusted. Correct key: Authentic.')
case(g,'V4','D1: protected cover-bit tampering','Recreate D1 so it is LAST CREATED; trusted public key.','Click Alter cover value; Verify modified copy at depth 1/manual 100. Then Restore original stego and verify again.','Modified copy: Tampered / COVER_HASH_MISMATCH; signature/content hash/decryption pass; cover check fails. Restored original: Authentic.')
case(g,'V5','D2: hidden encrypted-package tampering','Recreate D2 so it is LAST CREATED; trusted public key.','Click Alter hidden content; Verify modified copy at depth 2/manual 100. Restore original stego and verify again.','Modified copy: Cannot Verify / DECRYPTION_FAILED; no recovered content because AES-GCM fails. This is not expected to be Tampered. Restored original: Authentic.')
g=group('D. Three focused Compare checks','The positive comparisons are already part of D1–D4; only these additional edge cases are needed.')
case(g,'C1','Wrong image original: dimension mismatch','Original field: samples/covers/chelsea.png; stego field: saved D1 stego PNG.','Compare objects; then replace original with astronaut.png and compare again.','Wrong pair rejected: Objects must have matching dimensions/channels or audio format/duration. Correct original: successful D1 comparison.')
case(g,'C2','Wrong audio original: format/duration mismatch','Original field: samples/audio/sample-002.wav; stego field: saved D2 stego WAV.','Compare objects; then replace original with sample-006.wav and compare again.','Wrong pair rejected with same matching-properties message; rates/frame counts differ. Correct original: successful audio comparison.')
case(g,'C3','Identical-object comparison sanity check','samples/covers/astronaut.png in BOTH fields; optionally repeat sample-006.wav in both fields.','Compare identical objects.','Changed values=0, changed%=0, MSE=0, max difference=0; PSNR shows Identical. Optional audio: SNR inf (stego is identical to the cover). No authenticity verdict is inferred.')
end="""Completion checklist: four baseline flows done □; capacity/settings checks done □; all five negative cases produce non-Authentic and restore to Authentic □; three Compare checks done □. Attach capacity/result screenshots and recovered files/reports for D1–D4; attach actual error/verdict text for negatives.

Defect record: test ID; expected vs observed result; steps; screenshot/report; severity; retest. A mismatch/error is a successful negative test when it is the intended outcome. Do not count an API-only reference check as a manual browser result.

Reference scope: expected results were derived from the current frontend and backend implementation. The earlier reference_checks.json covers a broader controlled backend matrix; it is not evidence that all 16 rows in this reduced manual template have been executed. Keep your observation fields blank until you perform them.
"""
for i,g in enumerate(sections):
    g['title']=chr(65+i)+'. '+g['title'].split('. ',1)[1]
count=sum(len(g['cases']) for g in sections)
header=f'# Protect, Verify and Compare — focused four-demo test template\n\nSource commit: {revision}; prepared {datetime.now(timezone.utc).date()}; {count} independently numbered cases.\n\n'
md=header+'## Setup and execution record\n\n'+setup+'\n## Expected behavior and limits\n\n'+behavior+'\n## Common phase procedure\n\n'+workflow+'\n## File inventory\n\n| Alias | File(s), relative to repo root | Suitability |\n|---|---|---|\n'
for a,b,c in inventory:md+=f'| {a} | {b} | {c} |\n'
for g in sections:
 md+='\n## '+g['title']+'\n\n'+g['intro']+'\n\n'
 md+='| ID / test description | Files / baseline | Procedure / settings | Expected outcome | Observed result / evidence / Pass-Fail |\n|---|---|---|---|---|\n'
 for c in g['cases']:
  vals=[c['id']+' — '+c['description'],c['files'],c['steps'],c['expected'],'Protect: ____; Verify: ____; Compare: ____; Status: ____; Evidence/defect: ____']
  md+='| '+' | '.join(v.replace('|','/').replace('\n','<br>') for v in vals)+' |\n'
md+='\n## Evidence, defect log and reference scope\n\n'+end+'\n| Defect ID | Test ID | Expected | Actual / reproduction | Severity | Evidence | Retest |\n|---|---|---|---|---|---|---|\n| ____ | ____ | ____ | ____ | ____ | ____ | ____ |\n'
(OUT/'phase_test_template.md').write_text(md)
(OUT/'test_catalog.json').write_text(json.dumps({'source_commit':revision,'case_count':count,'sections':sections},indent=2)+'\n')

from docx import Document
from docx.shared import Inches,Pt
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc=Document();section=doc.sections[0];section.orientation=WD_ORIENT.LANDSCAPE
section.page_width=Inches(11.69);section.page_height=Inches(8.27)
section.top_margin=section.bottom_margin=Inches(.5);section.left_margin=section.right_margin=Inches(.45)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
normal.paragraph_format.space_after=Pt(5)
doc.add_heading('Protect, Verify and Compare',0)
doc.add_paragraph('Focused four-demo test template',style='Subtitle')
doc.add_paragraph(f'Source commit {revision} • {count} test cases • Prepared {datetime.now(timezone.utc).date()}')
for title,content in [('Setup and execution record',setup),('Expected behavior and limits',behavior),('Common phase procedure',workflow)]:
 doc.add_heading(title,1)
 for p in content.split('\n\n'):doc.add_paragraph(p)
doc.add_heading('File inventory — paths relative to repository root',1)
t=doc.add_table(rows=1,cols=3);t.style='Table Grid'
for cell,text in zip(t.rows[0].cells,['Alias','File(s)','Suitability']):cell.text=text
for row in inventory:
 for cell,text in zip(t.add_row().cells,row):cell.text=text
for g in sections:
 doc.add_page_break();doc.add_heading(g['title'],1);doc.add_paragraph(g['intro'])
 t=doc.add_table(rows=1,cols=5);t.style='Table Grid';t.autofit=False
 widths=[1.15,1.9,2.05,3.15,2.4]
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for cell,text,w in zip(t.rows[0].cells,['ID / description','Files / baseline','Procedure / settings','Expected outcome','Observed result / status / evidence'],widths):cell.text=text;cell.width=Inches(w)
 repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for c in g['cases']:
  cells=t.add_row().cells
  vals=[c['id']+'\n'+c['description'],c['files'],c['steps'],c['expected'],'Protect: __________\nVerify: __________\nCompare: __________\nStatus: __________\nEvidence/defect: __________\nNotes: __________']
  for cell,text,w in zip(cells,vals,widths):cell.text=text;cell.width=Inches(w)
 for row in t.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(3)
    for run in p.runs:run.font.size=Pt(9)
doc.add_page_break();doc.add_heading('Evidence, defect log and reference scope',1)
for p in end.split('\n\n'):doc.add_paragraph(p)
t=doc.add_table(rows=1,cols=7);t.style='Table Grid'
for cell,text in zip(t.rows[0].cells,['Defect ID','Test ID','Expected','Actual / steps','Severity','Evidence','Retest']):cell.text=text
for _ in range(8):
 for cell in t.add_row().cells:cell.text='________________'
footer=section.footer.paragraphs[0];footer.text='Stego Integrity Verifier • Manual testing • '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.save(OUT/'phase_test_template.docx')
print('Generated',count,'cases in Word and Markdown.')
